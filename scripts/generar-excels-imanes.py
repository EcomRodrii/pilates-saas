#!/usr/bin/env python3
"""
Genera los dos Excel que se descargan a cambio del email desde /recursos:

  · modelo-estudio-12-meses.xlsx      (guía «¿Es rentable un estudio de pilates?»)
  · simulador-escalera-bonos.xlsx     (guía «Bonos de pilates»)

Completan las calculadoras que ya hay en esas guías, no las sustituyen: la
calculadora da una cifra; el Excel deja jugar con los números mes a mes.

Las fórmulas son fórmulas de Excel de verdad (nada de valores pegados) y se
verifican con `scripts/verificar-excels-imanes.py`, que las calcula con un
motor independiente y las compara con una cuenta hecha aparte.

  python3 -m venv /tmp/venv-xl && /tmp/venv-xl/bin/pip install openpyxl formulas
  /tmp/venv-xl/bin/python scripts/generar-excels-imanes.py /tmp/imanes
  /tmp/venv-xl/bin/python scripts/verificar-excels-imanes.py /tmp/imanes

Después: copiarlos a public/recursos/descargas/ con los 10 primeros caracteres
de su sha256 en el nombre (`shasum -a 256`) y cambiar `archivo` en
lib/recursos/descargas.ts. El hash en el nombre evita cachés viejas.

Convención de colores (la misma que la plantilla de asistencia): cabecera
oliva = título; celda crema con letra azul = la rellenas tú; celda gris =
fórmula, no escribas encima.
"""
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OLIVA = "343825"
CREMA = "FFF6D6"
GRIS = "EFEFE9"
AZUL = "1F4E9A"
FUENTE = "Arial"

SEMANAS_MES = 4.33  # 52 ÷ 12, el mismo número que la guía de rentabilidad
IVA = 0.21          # IVA de una clase en un estudio privado (guía de IVA)

fino = Side(style="thin", color="CFCFC6")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)

EUR = '#,##0.00 "€"'
EUR0 = '#,##0 "€"'
PCT = "0.0%"
NUM1 = "#,##0.0"


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FUENTE, bold=bold, color=color, size=size, italic=italic)


def titulo(ws, celda, texto, ancho=None):
    ws[celda] = texto
    ws[celda].font = f(True, "FFFFFF", 12)
    ws[celda].fill = PatternFill("solid", fgColor=OLIVA)
    ws[celda].alignment = Alignment(vertical="center")


def cabecera(ws, fila, textos, col0=1):
    for i, t in enumerate(textos):
        c = ws.cell(row=fila, column=col0 + i, value=t)
        c.font = f(True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=OLIVA)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE


def entrada(ws, celda, valor, formato=None):
    c = ws[celda]
    c.value = valor
    c.font = f(False, AZUL)
    c.fill = PatternFill("solid", fgColor=CREMA)
    c.border = BORDE
    if formato:
        c.number_format = formato
    return c


def calculo(ws, celda, formula, formato=None, negrita=False):
    c = ws[celda]
    c.value = formula
    c.font = f(negrita)
    c.fill = PatternFill("solid", fgColor=GRIS)
    c.border = BORDE
    if formato:
        c.number_format = formato
    return c


def etiqueta(ws, celda, texto, negrita=False):
    c = ws[celda]
    c.value = texto
    c.font = f(negrita)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    return c


def nota(ws, celda, texto):
    c = ws[celda]
    c.value = texto
    c.font = f(False, "5A5A52", 9, italic=True)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    return c


def anchos(ws, valores):
    for i, w in enumerate(valores, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def leeme(wb, titulo_txt, parrafos):
    ws = wb.active
    ws.title = "Léeme"
    ws.sheet_view.showGridLines = False
    anchos(ws, [110])
    titulo(ws, "A1", titulo_txt)
    ws.row_dimensions[1].height = 26
    fila = 3
    for p in parrafos:
        c = ws.cell(row=fila, column=1, value=p)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[fila].height = max(18, 15 * (len(p) // 105 + 1))
        fila += 1
    # leyenda de colores
    fila += 1
    c = ws.cell(row=fila, column=1, value="Cómo leer los colores")
    c.font = f(True)
    fila += 1
    for texto, color, letra in [
        ("Celda crema con letra azul: la rellenas tú con tus números.", CREMA, AZUL),
        ("Celda gris: es una fórmula. No escribas encima: si lo haces, borras la fórmula y la hoja deja de cuadrar.", GRIS, "000000"),
    ]:
        c = ws.cell(row=fila, column=1, value=texto)
        c.font = f(False, letra)
        c.fill = PatternFill("solid", fgColor=color)
        c.border = BORDE
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[fila].height = 30
        fila += 1


# ─────────────────────────────────────────────────────────────────────────────
# 1. Modelo de tu estudio a 12 meses
# ─────────────────────────────────────────────────────────────────────────────

def modelo_12_meses(destino: Path):
    wb = Workbook()
    leeme(wb, "Modelo de tu estudio de pilates a 12 meses", [
        "Qué hace: con los números de tu estudio, te dice mes a mes cuánto ingresas, cuánto gastas, cuánto te queda y a partir de qué ocupación cubres costes y tu propio sueldo. Es la cuenta de la guía «¿Es rentable un estudio de pilates?» llevada a un año.",
        "Cómo se usa: 1) En la hoja «Supuestos» cambia las celdas crema por tus números. 2) Mira «12 meses» para ver el año. 3) Mira «Equilibrio» para saber la ocupación que necesitas. Prueba a subir el precio o a quitar una clase y vuelve a mirar.",
        "La fórmula de ingresos: plazas por clase × clases a la semana × 4,33 × ocupación × precio medio por plaza sin IVA. 4,33 son las semanas de un mes medio (52 ÷ 12).",
        "Los números que trae son un EJEMPLO para ver cómo funciona, no una recomendación ni la media del sector. Sustitúyelos por los tuyos.",
        "Todos los importes van SIN IVA. Si tus precios publicados llevan IVA, divídelos entre 1,21 antes de ponerlos (el IVA de una clase en un estudio privado es el 21 %).",
        "No es asesoramiento fiscal ni financiero: para impuestos, amortizaciones o una financiación, consúltalo con tu asesoría.",
    ])

    s = wb.create_sheet("Supuestos")
    s.sheet_view.showGridLines = False
    anchos(s, [46, 16, 70])
    titulo(s, "A1", "Supuestos: tus números")
    s.merge_cells("A1:C1")
    cabecera(s, 3, ["Concepto", "Tu valor", "Qué poner"])
    filas = [
        ("Plazas por clase (aforo)", 8, "0", "Las camas o plazas de una clase. Si varía, pon la media."),
        ("Clases a la semana", 30, "0", "Todas las clases que das en la semana, de todos los tipos."),
        ("Ocupación del primer mes", 0.45, PCT, "Qué parte de las plazas se llena el primer mes (0,45 = 45 %)."),
        ("Cuánto sube la ocupación cada mes (puntos)", 0.03, PCT, "Puntos de ocupación que ganas cada mes (0,03 = 3 puntos)."),
        ("Ocupación máxima que esperas", 0.8, PCT, "El techo realista de tu estudio. Más de 85 % cuesta sostenerlo."),
        ("Precio medio por plaza, SIN IVA (€)", 17, EUR, "Lo que ingresas de media por una plaza ocupada, quitado el IVA."),
        ("Comisión de cobro (%)", 0.015, PCT, "Lo que se come el cobro con tarjeta o domiciliación (pon 0 si no hay)."),
        ("Coste de la instructora por clase (€)", 22, EUR, "Lo que te cuesta una clase en instructora, con su Seguridad Social."),
        ("Alquiler del local al mes (€)", 1800, EUR, "Sin IVA."),
        ("Suministros y gastos generales al mes (€)", 350, EUR, "Luz, agua, limpieza, internet…"),
        ("Software, seguros y asesoría al mes (€)", 300, EUR, "Programa de reservas, seguro de responsabilidad, gestoría."),
        ("Marketing al mes (€)", 250, EUR, "Anuncios, redes, material."),
        ("Reserva para máquinas y reposición al mes (€)", 200, EUR, "Lo que apartas cada mes para mantener y reponer máquinas."),
        ("Lo que quieres sacar tú al mes (€)", 1800, EUR, "Tu sueldo objetivo. Si dejas 0, el modelo solo mira si cubres costes."),
        ("Inversión inicial por recuperar (€)", 30000, EUR0, "Opcional: máquinas, obra, fianza. Pon 0 si no aplica."),
    ]
    for i, (conc, val, fmt, ayuda) in enumerate(filas, start=4):
        etiqueta(s, f"A{i}", conc)
        entrada(s, f"B{i}", val, fmt)
        nota(s, f"C{i}", ayuda)
        s.row_dimensions[i].height = 28
    nota(s, "A20", "Semanas de un mes medio (fijo)")
    calculo(s, "B20", SEMANAS_MES, "0.00")
    nota(s, "C20", "52 ÷ 12 = 4,33. No lo cambies.")

    # nombres cortos de las celdas de entrada
    PLAZAS, CLASES, OCU0, PASO, OCUMAX, PRECIO, COM, INSTR, ALQ, SUM, SOFT, MKT, RES, SUELDO, INV, SEM = (
        "Supuestos!$B$4", "Supuestos!$B$5", "Supuestos!$B$6", "Supuestos!$B$7", "Supuestos!$B$8", "Supuestos!$B$9",
        "Supuestos!$B$10", "Supuestos!$B$11", "Supuestos!$B$12", "Supuestos!$B$13", "Supuestos!$B$14",
        "Supuestos!$B$15", "Supuestos!$B$16", "Supuestos!$B$17", "Supuestos!$B$18", "Supuestos!$B$20",
    )

    m = wb.create_sheet("12 meses")
    m.sheet_view.showGridLines = False
    anchos(m, [40] + [12] * 12 + [14])
    titulo(m, "A1", "Tu año, mes a mes (todo sin IVA)")
    m.merge_cells("A1:N1")
    cabecera(m, 3, ["Concepto"] + [f"Mes {i}" for i in range(1, 13)] + ["Total año"])
    filas_m = [
        # (fila, etiqueta, fórmula(col) , formato, total)
        (4, "Mes", lambda c, i: i, "0", None),
        (5, "Ocupación", lambda c, i: f"=MIN({OCUMAX},{OCU0}+({c}4-1)*{PASO})", PCT, None),
        (6, "Plazas ofrecidas", lambda c, i: f"={PLAZAS}*{CLASES}*{SEM}", NUM1, "SUM"),
        (7, "Plazas ocupadas", lambda c, i: f"={c}6*{c}5", NUM1, "SUM"),
        (8, "Ingresos", lambda c, i: f"={c}7*{PRECIO}", EUR0, "SUM"),
        (9, "Comisiones de cobro", lambda c, i: f"={c}8*{COM}", EUR0, "SUM"),
        (10, "Instructoras", lambda c, i: f"={INSTR}*{CLASES}*{SEM}", EUR0, "SUM"),
        (11, "Costes fijos (local, gastos, software, marketing, reserva)", lambda c, i: f"={ALQ}+{SUM}+{SOFT}+{MKT}+{RES}", EUR0, "SUM"),
        (12, "Resultado antes de ti", lambda c, i: f"={c}8-{c}9-{c}10-{c}11", EUR0, "SUM"),
        (13, "Tu sueldo objetivo", lambda c, i: f"={SUELDO}", EUR0, "SUM"),
        (14, "Resultado después de tu sueldo", lambda c, i: f"={c}12-{c}13", EUR0, "SUM"),
        (15, "Acumulado antes de ti", None, EUR0, None),
        (16, "Acumulado después de tu sueldo", None, EUR0, None),
        (17, "¿Cubre costes?", lambda c, i: f'=IF({c}12>=0,"Sí","No")', None, None),
        (18, "¿Cubre costes y tu sueldo?", lambda c, i: f'=IF({c}14>=0,"Sí","No")', None, None),
        (19, "Inversión recuperada (1 = sí)", lambda c, i: f"=IF({c}16>={INV},1,0)", "0", None),
    ]
    for fila, etiq, fn, fmt, tot in filas_m:
        etiqueta(m, f"A{fila}", etiq, negrita=fila in (12, 14))
        for i in range(1, 13):
            col = get_column_letter(1 + i)
            if fila == 15:
                formula = f"={col}12" if i == 1 else f"={get_column_letter(i)}15+{col}12"
            elif fila == 16:
                formula = f"={col}14" if i == 1 else f"={get_column_letter(i)}16+{col}14"
            else:
                formula = fn(col, i)
            calculo(m, f"{col}{fila}", formula, fmt, negrita=fila in (12, 14))
        if tot:
            calculo(m, f"N{fila}", f"=SUM(B{fila}:M{fila})", fmt, negrita=True)
        else:
            calculo(m, f"N{fila}", None)
    m.freeze_panes = "B4"

    e = wb.create_sheet("Equilibrio")
    e.sheet_view.showGridLines = False
    anchos(e, [58, 18, 62])
    titulo(e, "A1", "¿Qué ocupación necesitas?")
    e.merge_cells("A1:C1")
    cabecera(e, 3, ["Cuenta", "Resultado", "Cómo se calcula"])
    cuentas = [
        (4, "Plazas ofrecidas al mes", f"={PLAZAS}*{CLASES}*{SEM}", NUM1, "Plazas por clase × clases a la semana × 4,33."),
        (5, "Lo que deja cada plaza (€)", f"={PRECIO}*(1-{COM})", EUR, "Precio medio sin IVA menos la comisión de cobro."),
        (6, "Costes del mes sin tu sueldo (€)", f"={INSTR}*{CLASES}*{SEM}+{ALQ}+{SUM}+{SOFT}+{MKT}+{RES}", EUR0, "Instructoras + local + gastos + software + marketing + reserva."),
        (7, "Ocupación de equilibrio, sin tu sueldo", "=IFERROR(B6/(B4*B5),0)", PCT, "Costes del mes ÷ (plazas del mes × lo que deja cada plaza)."),
        (8, "Ocupación de equilibrio, con tu sueldo", f"=IFERROR((B6+{SUELDO})/(B4*B5),0)", PCT, "La misma cuenta sumando a los costes lo que quieres sacar tú. Esta es la que necesitas de verdad."),
        (9, "Personas por clase para cubrir costes", f"=B7*{PLAZAS}", NUM1, "Ocupación de equilibrio × plazas por clase."),
        (10, "Personas por clase para cubrir costes y tu sueldo", f"=B8*{PLAZAS}", NUM1, "Ocupación de equilibrio con sueldo × plazas por clase."),
        (11, "¿Llegas con tu ocupación máxima?", f'=IF({OCUMAX}>=B8,"Sí","No")', None, "Compara tu techo de ocupación con la ocupación de equilibrio con sueldo."),
        (12, "Precio mínimo por plaza para llegar con tu ocupación máxima (€)", f"=IFERROR((B6+{SUELDO})/(B4*{OCUMAX}*(1-{COM})),0)", EUR, "El precio medio sin IVA que necesitas para cubrir todo si llenas hasta tu techo."),
        (13, "Mes en que recuperas la inversión", '=IFERROR(INDEX(\'12 meses\'!B4:M4,MATCH(1,\'12 meses\'!B19:M19,0)),"No en 12 meses")', "0", "Primer mes con el acumulado después de tu sueldo por encima de la inversión inicial."),
    ]
    for fila, etiq, formula, fmt, como in cuentas:
        etiqueta(e, f"A{fila}", etiq, negrita=fila in (8, 12))
        calculo(e, f"B{fila}", formula, fmt, negrita=fila in (8, 12))
        nota(e, f"C{fila}", como)
        e.row_dimensions[fila].height = 30
    wb.save(destino)


# ─────────────────────────────────────────────────────────────────────────────
# 2. Simulador de la escalera de bonos
# ─────────────────────────────────────────────────────────────────────────────

def simulador_bonos(destino: Path):
    wb = Workbook()
    leeme(wb, "Simulador de tu escalera de precios de bonos y cuotas", [
        "Qué hace: parte del precio de tu clase suelta y de un descuento por escalón, te dice lo que sale cada sesión y el ahorro de cada producto, y te deja comparar hasta tres escenarios de precios con lo que vendes al mes. Es la escalera de la guía «Bonos de pilates» con tus números.",
        "Cómo se usa: 1) En «Escalera» pon tu clase suelta, tus descuentos y lo que te cuesta una plaza. 2) En «Escenarios» pon cuántos productos vendes al mes y prueba otros precios en las columnas «A» y «B». 3) Mira las alertas en rojo.",
        "Dos errores que avisa: un escalón por debajo de lo que te cuesta una plaza (vendes con pérdida), y un bono de 10 más barato por sesión que la cuota de una clase semanal (empujas a tus alumnas fuera de la cuota).",
        "Los precios llevan IVA (es lo que publican los estudios); el coste por plaza va SIN IVA. Para comparar, el modelo quita el 21 % al precio antes de mirar si cubre el coste.",
        "Los números que trae son un ejemplo (25 € la suelta, bonos al 12 % y al 20 %, cuotas al 24 % y 34 %). Sustitúyelos por los tuyos.",
        "Las cuotas se calculan con cuatro clases al mes (una a la semana) y ocho (dos a la semana). Es una aproximación: un mes real tiene entre 4 y 5 semanas.",
        "No es asesoramiento fiscal ni de precios: es una calculadora. Revisa con tu asesoría cualquier decisión con impacto fiscal.",
    ])

    s = wb.create_sheet("Escalera")
    s.sheet_view.showGridLines = False
    anchos(s, [34, 14, 16, 14, 14, 16, 22])
    titulo(s, "A1", "Tu escalera de precios")
    s.merge_cells("A1:G1")
    etiqueta(s, "A3", "Precio de la clase suelta, con IVA (€)")
    entrada(s, "B3", 25, EUR)
    etiqueta(s, "A4", "Lo que te cuesta una plaza ocupada, SIN IVA (€)")
    entrada(s, "B4", 11, EUR)
    nota(s, "C4", "Pon 0 si no quieres que avise del coste. Cómo calcularlo: guía «¿Es rentable un estudio de pilates?».")
    etiqueta(s, "A5", "IVA de una clase (fijo)")
    calculo(s, "B5", IVA, PCT)
    s.row_dimensions[4].height = 30

    cabecera(s, 7, ["Producto", "Sesiones", "Descuento por sesión", "Precio por sesión (€)", "Precio del producto (€)", "Ahorro frente a la suelta", "Alerta"])
    prods = [
        (8, "Clase suelta", 1, 0.0),
        (9, "Bono de 5 sesiones", 5, 0.12),
        (10, "Bono de 10 sesiones", 10, 0.20),
        (11, "Cuota de 1 clase a la semana (al mes)", 4, 0.24),
        (12, "Cuota de 2 clases a la semana (al mes)", 8, 0.34),
    ]
    for fila, nombre, ses, dto in prods:
        etiqueta(s, f"A{fila}", nombre)
        calculo(s, f"B{fila}", ses, "0")
        if fila == 8:
            calculo(s, f"C{fila}", 0, PCT)
        else:
            entrada(s, f"C{fila}", dto, PCT)
        calculo(s, f"D{fila}", f"=ROUND($B$3*(1-C{fila}),2)", EUR)
        calculo(s, f"E{fila}", f"=ROUND(D{fila}*B{fila},2)", EUR, negrita=True)
        calculo(s, f"F{fila}", f"=IF($B$3>0,ROUND((1-D{fila}/$B$3)*100,0)/100,0)", "0%")
        calculo(s, f"G{fila}", f'=IF(AND($B$4>0,D{fila}/(1+$B$5)<$B$4),"Por debajo de tu coste","")')
        s.row_dimensions[fila].height = 22
    etiqueta(s, "A14", "¿El bono de 10 sale más barato por sesión que la cuota de 1 clase semanal?", negrita=True)
    s.merge_cells("A14:E14")
    s.row_dimensions[14].height = 30
    calculo(s, "F14", '=IF($B$3*(1-C10)<$B$3*(1-C11),"Sí: empuja fuera de la cuota","No")', None, negrita=True)
    s.merge_cells("F14:G14")
    nota(s, "A16", "Las cuotas son el precio al mes; «sesiones» es cuántas clases incluye ese mes (4 y 8).")
    s.merge_cells("A16:G16")

    sc = wb.create_sheet("Escenarios")
    sc.sheet_view.showGridLines = False
    anchos(sc, [52, 18, 18, 18])
    titulo(sc, "A1", "Compara tres escenarios con lo que vendes al mes")
    sc.merge_cells("A1:D1")
    cabecera(sc, 3, ["", "Ahora", "Escenario A", "Escenario B"])
    etiqueta(sc, "A4", "Descuentos por sesión", negrita=True)
    desc = [
        (5, "Bono de 5 sesiones", "C9", [0.12, 0.10, 0.15]),
        (6, "Bono de 10 sesiones", "C10", [0.20, 0.18, 0.22]),
        (7, "Cuota de 1 clase a la semana", "C11", [0.24, 0.24, 0.26]),
        (8, "Cuota de 2 clases a la semana", "C12", [0.34, 0.34, 0.36]),
    ]
    for fila, nombre, ref, vals in desc:
        etiqueta(sc, f"A{fila}", nombre)
        calculo(sc, f"B{fila}", f"=Escalera!{ref}", PCT)
        entrada(sc, f"C{fila}", vals[1], PCT)
        entrada(sc, f"D{fila}", vals[2], PCT)
    nota(sc, "A9", "«Ahora» sale de la hoja «Escalera». En A y B pon los descuentos que quieres probar.")
    sc.merge_cells("A9:D9")
    etiqueta(sc, "A11", "Lo que vendes al mes (unidades)", negrita=True)
    ventas = [
        (12, "Clases sueltas", 60, 8),
        (13, "Bonos de 5 sesiones", 12, 9),
        (14, "Bonos de 10 sesiones", 10, 10),
        (15, "Cuotas de 1 clase a la semana", 25, 11),
        (16, "Cuotas de 2 clases a la semana", 15, 12),
    ]
    for fila, nombre, base, _ in ventas:
        etiqueta(sc, f"A{fila}", nombre)
        entrada(sc, f"B{fila}", base, "0")
        entrada(sc, f"C{fila}", base, "0")
        entrada(sc, f"D{fila}", base, "0")
    nota(sc, "A17", "Si un escenario cambia lo que vendes (más bonos, menos sueltas…), cámbialo aquí. Cada columna es independiente.")
    sc.merge_cells("A17:D17")

    etiqueta(sc, "A19", "Resultado del mes", negrita=True)
    # precio por sesión de cada escalón en cada escenario (fila auxiliar)
    sub = [
        (20, "Precio por sesión del bono de 5 (€)", lambda c: f"=ROUND(Escalera!$B$3*(1-{c}5),2)", EUR),
        (21, "Precio por sesión del bono de 10 (€)", lambda c: f"=ROUND(Escalera!$B$3*(1-{c}6),2)", EUR),
        (22, "Precio por sesión de la cuota de 1 clase (€)", lambda c: f"=ROUND(Escalera!$B$3*(1-{c}7),2)", EUR),
        (23, "Precio por sesión de la cuota de 2 clases (€)", lambda c: f"=ROUND(Escalera!$B$3*(1-{c}8),2)", EUR),
    ]
    for fila, etiq, fn, fmt in sub:
        etiqueta(sc, f"A{fila}", etiq)
        for c in "BCD":
            calculo(sc, f"{c}{fila}", fn(c), fmt)
    res = [
        (24, "Ingresos del mes, con IVA (€)", lambda c: (
            f"=ROUND({c}12*Escalera!$B$3+{c}13*{c}20*5+{c}14*{c}21*10+{c}15*{c}22*4+{c}16*{c}23*8,2)"), EUR0, True),
        (25, "Ingresos del mes, sin IVA (€)", lambda c: f"=ROUND({c}24/(1+Escalera!$B$5),2)", EUR0, True),
        (26, "Sesiones vendidas en el mes", lambda c: f"={c}12+{c}13*5+{c}14*10+{c}15*4+{c}16*8", "#,##0", False),
        (27, "Coste de esas plazas (€)", lambda c: f"=ROUND({c}26*Escalera!$B$4,2)", EUR0, False),
        (28, "Margen del mes sobre el coste de la plaza (€)", lambda c: f"=ROUND({c}25-{c}27,2)", EUR0, True),
        (29, "Ingreso medio por sesión, sin IVA (€)", lambda c: f"=IF({c}26>0,ROUND({c}25/{c}26,2),0)", EUR, False),
        (30, "Diferencia de margen frente a «Ahora» (€)", lambda c: f"=ROUND({c}28-$B$28,2)", EUR0, True),
        (31, "Alerta: algún escalón por debajo de tu coste",
         lambda c: (f'=IF(AND(Escalera!$B$4>0,MIN(Escalera!$D$8,{c}20,{c}21,{c}22,{c}23)/(1+Escalera!$B$5)<Escalera!$B$4),"Sí","No")'), None, False),
        (32, "Alerta: el bono de 10 canibaliza la cuota", lambda c: f'=IF({c}21<{c}22,"Sí","No")', None, False),
    ]
    for fila, etiq, fn, fmt, neg in res:
        etiqueta(sc, f"A{fila}", etiq, negrita=neg)
        for c in "BCD":
            calculo(sc, f"{c}{fila}", fn(c), fmt, negrita=neg)
        sc.row_dimensions[fila].height = 20
    nota(sc, "A33", "El margen es una aproximación: supone que cada sesión vendida se da en el mes y mide solo el coste de la plaza, no el resto de tus gastos fijos.")
    sc.merge_cells("A33:D33")
    sc.row_dimensions[33].height = 30
    wb.save(destino)


if __name__ == "__main__":
    salida = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/imanes")
    salida.mkdir(parents=True, exist_ok=True)
    modelo_12_meses(salida / "modelo-estudio-12-meses.xlsx")
    simulador_bonos(salida / "simulador-escalera-bonos.xlsx")
    print("Generados en", salida)
